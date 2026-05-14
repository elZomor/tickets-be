from django.contrib import admin, messages
from django.http import HttpResponse
from django.template.response import TemplateResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO

from hita_evaluation.models import (
    Professor,
    ProfessorGrade,
    Subject,
    Course,
    CourseProfessor,
    Regulation,
    Semester,
    SemesterType,
    SurveyTemplate,
    QuestionCategory,
    SurveyQuestion,
    SurveySession,
    SurveyAnswer,
    Department,
    Category,
    QuestionType,
)


def generate_evaluation_report(courses, filename_prefix='evaluation_report'):
    """Generate Excel report for given courses with all answers."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Evaluation Report'

    # Define styles
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # Headers
    headers = [
        'الفصل الدراسي',  # Semester
        'اللائحة',  # Regulation
        'القسم',  # Department
        'المادة',  # Subject
        'التصنيف',  # Category
        'الساعات',  # Credit Hours
        'الأستاذ',  # Professor
        'درجة الأستاذ',  # Professor Grade
        'تصنيف السؤال',  # Question Category
        'السؤال',  # Question
        'نوع السؤال',  # Question Type
        'الإجابة',  # Answer
        'تاريخ التقييم',  # Evaluation Date
        'نوع التعليم',  # Education Type (Parallel/Credit)
    ]

    # Write headers
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Get all answers for the given courses
    answers = SurveyAnswer.objects.filter(course__in=courses).select_related(
        'survey_session',
        'survey_session__regulation',
        'course',
        'course__subject',
        'course__semester',
        'course__regulations',
        'professor',
        'question',
        'question__question_category',
    ).order_by(
        'course__semester__year',
        'course__subject__department',
        'course__subject__name',
        'professor__full_name',
        'question__question_category__name',
        'question__id',
    )

    row = 2

    # If no answers, add a message row
    if not answers.exists():
        cell = ws.cell(row=2, column=1, value='لا توجد تقييمات لهذه المقررات')
        cell.alignment = Alignment(horizontal='center')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))

    for answer in answers:
        course = answer.course
        subject = course.subject
        semester = course.semester
        professor = answer.professor
        question = answer.question
        session = answer.survey_session

        # Determine answer value based on question type
        if question.question_type == 'YN':
            answer_value = 'نعم' if answer.yes_no_answer else 'لا' if answer.yes_no_answer is False else '-'
        elif question.question_type in ['R', 'S']:
            answer_value = str(answer.rating_answer) if answer.rating_answer else '-'
        elif question.question_type == 'T':
            answer_value = answer.text_answer or '-'
        else:
            answer_value = '-'

        # Semester display
        semester_display = f"{semester.year} - {semester.get_type_display()}"

        # Regulation (from session or course)
        regulation_name = '-'
        if session.regulation:
            regulation_name = session.regulation.name
        elif course.regulations:
            regulation_name = course.regulations.name

        # Professor grade from CourseProfessor through model
        professor_grade = '-'
        try:
            course_professor = CourseProfessor.objects.get(
                course=course, professor=professor
            )
            if course_professor.grade:
                professor_grade = course_professor.get_grade_display()
        except CourseProfessor.DoesNotExist:
            pass

        # Education type
        education_type = 'التعليم الموازي' if session.is_parallel else 'الساعات المعتمدة'

        row_data = [
            semester_display,
            regulation_name,
            subject.get_department_display() if subject.department else '-',
            subject.name,
            subject.get_category_display(),
            subject.credit_hours,
            professor.full_name,
            professor_grade,
            question.question_category.name if question.question_category else '-',
            question.question_text,
            question.get_question_type_display(),
            answer_value,
            session.created_at.strftime('%Y-%m-%d %H:%M'),
            education_type,
        ]

        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

        row += 1

    # Adjust column widths
    column_widths = [15, 15, 25, 25, 20, 10, 20, 18, 20, 50, 15, 15, 18, 18]
    for col, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    # Set RTL for the sheet
    ws.sheet_view.rightToLeft = True

    # Create response
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.xlsx"'

    return response


@admin.register(Regulation)
class RegulationAdmin(admin.ModelAdmin):
    list_display = ['name', 'valid_from', 'valid_to', 'is_latest']
    list_filter = ['is_latest']
    search_fields = ['name']


@admin.register(Professor)
class ProfessorAdmin(admin.ModelAdmin):
    list_display = ['full_name', 'department', 'grade']
    list_filter = ['department', 'grade']
    search_fields = ['full_name']


@admin.register(Subject)
class SubjectAdmin(admin.ModelAdmin):
    list_display = ['name', 'department', 'category', 'credit_hours']
    list_filter = ['department', 'category']
    search_fields = ['name']


@admin.register(Semester)
class SemesterAdmin(admin.ModelAdmin):
    list_display = ['year', 'type', 'is_current']
    list_filter = ['type', 'is_current', 'year']
    actions = ['export_department_report']
    change_list_template = 'admin/semester_changelist.html'

    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path(
                'import-semester/',
                self.admin_site.admin_view(self.import_semester_view),
                name='hita_evaluation_semester_import',
            ),
        ]
        return custom_urls + urls

    def import_semester_view(self, request):
        """Import semester data from Excel file."""
        from hita_evaluation.models import SemesterType, Category

        if request.method == 'POST' and request.FILES.get('excel_file'):
            excel_file = request.FILES['excel_file']

            try:
                wb = load_workbook(excel_file)
                ws = wb.active

                # Get headers from first row
                headers = [cell.value for cell in ws[1]]
                expected_headers = [
                    'year', 'type', 'department', 'subject_name',
                    'category', 'credit_hours', 'professor', 'is_split',
                    'regulation', 'is_parallel', 'professor_grade',
                ]
                required_headers = [
                    'year', 'type', 'department', 'subject_name',
                    'category', 'credit_hours', 'professor', 'is_split',
                ]

                # Validate required headers
                if not all(h in headers for h in required_headers):
                    self.message_user(
                        request,
                        f'الأعمدة المطلوبة: {", ".join(required_headers)}',
                        messages.ERROR,
                    )
                    return TemplateResponse(
                        request,
                        'admin/import_semester.html',
                        {
                            **self.admin_site.each_context(request),
                            'opts': self.model._meta,
                            'semester_types': [(st.value, st.label) for st in SemesterType],
                            'departments': [(d.value, d.label) for d in Department],
                            'categories': [(c.value, c.label) for c in Category],
                            'regulations': Regulation.objects.all(),
                            'professor_grades': [(g.value, g.label) for g in ProfessorGrade],
                        },
                    )

                # Get column indices (only for headers that exist)
                col_idx = {h: headers.index(h) for h in headers if h in expected_headers}

                # Get latest active survey template
                latest_template = SurveyTemplate.objects.filter(is_active=True).order_by('-version').first()

                # Track created objects
                semester = None
                subjects_created = 0
                professors_created = 0
                courses_created = 0

                valid_semester_types = [st.value for st in SemesterType]
                valid_departments = [d.value for d in Department]
                valid_categories = [c.value for c in Category]

                valid_professor_grades = [g.value for g in ProfessorGrade]

                for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):  # Skip empty rows
                        continue

                    year = row[col_idx['year']]
                    semester_type = row[col_idx['type']]
                    department = row[col_idx['department']]
                    subject_name = row[col_idx['subject_name']]
                    category = row[col_idx['category']]
                    credit_hours = row[col_idx['credit_hours']]
                    professor_str = row[col_idx['professor']]
                    is_split = row[col_idx['is_split']]

                    # Optional new columns
                    regulation_name = row[col_idx['regulation']] if 'regulation' in col_idx else None
                    is_parallel_val = row[col_idx['is_parallel']] if 'is_parallel' in col_idx else None
                    professor_grade_str = row[col_idx['professor_grade']] if 'professor_grade' in col_idx else None

                    # Skip rows without subject name
                    if not subject_name:
                        continue

                    # Create/get semester (use first row's data)
                    if semester is None:
                        if not year or not semester_type:
                            self.message_user(
                                request,
                                'السنة ونوع الفصل مطلوبان في الصف الأول',
                                messages.ERROR,
                            )
                            return TemplateResponse(
                                request,
                                'admin/import_semester.html',
                                {
                                    **self.admin_site.each_context(request),
                                    'opts': self.model._meta,
                                    'semester_types': [(st.value, st.label) for st in SemesterType],
                                    'departments': [(d.value, d.label) for d in Department],
                                    'categories': [(c.value, c.label) for c in Category],
                                },
                            )

                        if semester_type not in valid_semester_types:
                            self.message_user(
                                request,
                                f'نوع الفصل غير صالح: {semester_type}. الأنواع المتاحة: {", ".join(valid_semester_types)}',
                                messages.ERROR,
                            )
                            return TemplateResponse(
                                request,
                                'admin/import_semester.html',
                                {
                                    **self.admin_site.each_context(request),
                                    'opts': self.model._meta,
                                    'semester_types': [(st.value, st.label) for st in SemesterType],
                                    'departments': [(d.value, d.label) for d in Department],
                                    'categories': [(c.value, c.label) for c in Category],
                                },
                            )

                        semester, _ = Semester.objects.get_or_create(
                            year=int(year),
                            type=semester_type,
                            defaults={'is_current': True}
                        )
                        # Update is_current if semester exists
                        if not semester.is_current:
                            semester.is_current = True
                            semester.save()

                    # Validate department
                    if department and department not in valid_departments:
                        self.message_user(
                            request,
                            f'القسم غير صالح في الصف {row_num}: {department}',
                            messages.WARNING,
                        )
                        continue

                    # Validate category
                    if category and category not in valid_categories:
                        self.message_user(
                            request,
                            f'التصنيف غير صالح في الصف {row_num}: {category}',
                            messages.WARNING,
                        )
                        continue

                    # Create/get subject
                    subject, created = Subject.objects.get_or_create(
                        name=subject_name,
                        department=department,
                        defaults={
                            'category': category or Category.MANDATORY_DEPARTMENT,
                            'credit_hours': int(credit_hours) if credit_hours else 3,
                        }
                    )
                    if created:
                        subjects_created += 1

                    # Parse professors and their grades
                    professors = []
                    professor_grades = []
                    if professor_str:
                        professor_names = [p.strip() for p in str(professor_str).split('-') if p.strip()]
                        for prof_name in professor_names:
                            prof, created = Professor.objects.get_or_create(
                                full_name=prof_name,
                                defaults={'department': department}
                            )
                            professors.append(prof)
                            if created:
                                professors_created += 1

                    # Parse professor grades (same format as professors, using -)
                    if professor_grade_str:
                        grade_list = [g.strip() for g in str(professor_grade_str).split('-')]
                        for grade in grade_list:
                            if grade in valid_professor_grades:
                                professor_grades.append(grade)
                            else:
                                professor_grades.append('')
                    # Pad grades list to match professors
                    while len(professor_grades) < len(professors):
                        professor_grades.append('')

                    # Attach grade to professor profiles when provided
                    for idx, prof in enumerate(professors):
                        grade_value = professor_grades[idx] if idx < len(professor_grades) else ''
                        if grade_value and grade_value in valid_professor_grades and prof.grade != grade_value:
                            prof.grade = grade_value
                            prof.save(update_fields=['grade'])

                    # Parse is_split
                    if isinstance(is_split, bool):
                        split = is_split
                    elif isinstance(is_split, str):
                        split = is_split.lower() in ['true', 'yes', '1', 'نعم']
                    else:
                        split = bool(is_split) if is_split is not None else False

                    # Parse is_parallel
                    if isinstance(is_parallel_val, bool):
                        is_parallel = is_parallel_val
                    elif isinstance(is_parallel_val, str):
                        is_parallel = is_parallel_val.lower() in ['true', 'yes', '1', 'نعم']
                    else:
                        is_parallel = bool(is_parallel_val) if is_parallel_val is not None else False

                    # Get/create regulation
                    regulation = None
                    if regulation_name:
                        regulation, _ = Regulation.objects.get_or_create(
                            name=str(regulation_name).strip()
                        )

                    # Create course(s)
                    if split and len(professors) > 1:
                        # Create separate course for each professor
                        for idx, prof in enumerate(professors):
                            course = Course.objects.create(
                                subject=subject,
                                semester=semester,
                                survey_template=latest_template,
                                is_parallel=is_parallel,
                                regulations=regulation,
                            )
                            # Use through model to set professor with grade
                            CourseProfessor.objects.create(
                                course=course,
                                professor=prof,
                                grade=professor_grades[idx] if idx < len(professor_grades) else '',
                            )
                            courses_created += 1
                    else:
                        # Create one course with all professors
                        course = Course.objects.create(
                            subject=subject,
                            semester=semester,
                            survey_template=latest_template,
                            is_parallel=is_parallel,
                            regulations=regulation,
                        )
                        for idx, prof in enumerate(professors):
                            # Use through model to set professor with grade
                            CourseProfessor.objects.create(
                                course=course,
                                professor=prof,
                                grade=professor_grades[idx] if idx < len(professor_grades) else '',
                            )
                        courses_created += 1

                self.message_user(
                    request,
                    f'تم استيراد الفصل الدراسي "{semester}" بنجاح. '
                    f'المواد: {subjects_created}، الأساتذة: {professors_created}، المقررات: {courses_created}',
                    messages.SUCCESS,
                )
                from django.urls import reverse
                from django.http import HttpResponseRedirect
                return HttpResponseRedirect(reverse('admin:hita_evaluation_semester_changelist'))

            except Exception as e:
                self.message_user(
                    request,
                    f'خطأ في استيراد الملف: {str(e)}',
                    messages.ERROR,
                )

        from hita_evaluation.models import SemesterType, Category
        context = {
            **self.admin_site.each_context(request),
            'title': 'استيراد فصل دراسي - Import Semester',
            'opts': self.model._meta,
            'semester_types': [(st.value, st.label) for st in SemesterType],
            'departments': [(d.value, d.label) for d in Department],
            'categories': [(c.value, c.label) for c in Category],
            'regulations': Regulation.objects.all(),
            'professor_grades': [(g.value, g.label) for g in ProfessorGrade],
        }

        return TemplateResponse(
            request,
            'admin/import_semester.html',
            context,
        )

    @admin.action(description='تصدير تقرير القسم (Export Department Report)')
    def export_department_report(self, request, queryset):
        """Export report with department selection using intermediate page."""
        # If department is selected (POST with apply), generate report
        if 'apply' in request.POST:
            department = request.POST.get('department')
            if department:
                courses = Course.objects.filter(
                    semester__in=queryset,
                    subject__department=department,
                )

                if not courses.exists():
                    self.message_user(
                        request,
                        'لا توجد مقررات لهذا القسم في الفصول المحددة',
                        messages.WARNING,
                    )
                    return None

                semester_names = '_'.join([f"{s.year}_{s.type}" for s in queryset[:2]])
                dept_label = dict(Department.choices).get(department, department)
                filename = f'evaluation_{dept_label}_{semester_names}'

                return generate_evaluation_report(courses, filename)

        # Show intermediate page with department selection
        departments_choices = [
            (value, label)
            for value, label in Department.choices
            if value != Department.GENERAL
        ]

        context = {
            **self.admin_site.each_context(request),
            'title': 'تصدير تقرير التقييم - Export Evaluation Report',
            'queryset': queryset,
            'opts': self.model._meta,
            'action_checkbox_name': admin.helpers.ACTION_CHECKBOX_NAME,
            'departments': departments_choices,
            'media': self.media,
        }

        return TemplateResponse(
            request,
            'admin/department_selection.html',
            context,
        )


class CourseProfessorInline(admin.TabularInline):
    model = CourseProfessor
    extra = 1
    autocomplete_fields = ['professor']


@admin.register(Course)
class CourseAdmin(admin.ModelAdmin):
    list_display = ['subject', 'semester', 'survey_template', 'get_professors']
    list_filter = ['semester', 'subject__department', 'survey_template']
    search_fields = ['subject__name', 'professor__full_name']
    inlines = [CourseProfessorInline]
    actions = ['export_evaluation_report']

    def get_professors(self, obj):
        return ' - '.join([p.full_name for p in obj.professor.all()])

    get_professors.short_description = 'Professors'

    @admin.action(description='تصدير تقرير التقييم (Export Evaluation Report)')
    def export_evaluation_report(self, request, queryset):
        """Export evaluation report for selected courses."""
        course_names = '_'.join([c.subject.name[:10] for c in queryset[:3]])
        filename = f'evaluation_report_{course_names}'
        return generate_evaluation_report(queryset, filename)


@admin.register(SurveyTemplate)
class SurveyTemplateAdmin(admin.ModelAdmin):
    list_display = ['name', 'version', 'is_active', 'created_at', 'updated_at']
    list_filter = ['is_active']
    search_fields = ['name']
    change_list_template = 'admin/survey_template_changelist.html'

    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path(
                'import-template/',
                self.admin_site.admin_view(self.import_template_view),
                name='hita_evaluation_surveytemplate_import',
            ),
        ]
        return custom_urls + urls

    def import_template_view(self, request):
        """Import survey template from Excel file."""
        if request.method == 'POST' and request.FILES.get('excel_file'):
            excel_file = request.FILES['excel_file']

            try:
                wb = load_workbook(excel_file)
                ws = wb.active

                # Get headers from first row
                headers = [cell.value for cell in ws[1]]
                expected_headers = [
                    'survey_name', 'version', 'question_category',
                    'question_text', 'question_type', 'is_mandatory'
                ]
                optional_headers = ['question_sub_text']

                # Validate headers
                if not all(h in headers for h in expected_headers):
                    self.message_user(
                        request,
                        f'الأعمدة المطلوبة: {", ".join(expected_headers)}',
                        messages.ERROR,
                    )
                    return TemplateResponse(
                        request,
                        'admin/import_template.html',
                        {**self.admin_site.each_context(request), 'opts': self.model._meta},
                    )

                # Get column indices (required + optional headers that exist)
                all_headers = expected_headers + optional_headers
                col_idx = {h: headers.index(h) for h in all_headers if h in headers}

                # Track created objects
                template = None
                questions_created = 0
                categories_created = 0

                valid_question_types = [qt.value for qt in QuestionType]

                for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):  # Skip empty rows
                        continue

                    survey_name = row[col_idx['survey_name']]
                    version = row[col_idx['version']]
                    category_name = row[col_idx['question_category']]
                    question_text = row[col_idx['question_text']]
                    question_type = row[col_idx['question_type']]
                    is_mandatory = row[col_idx['is_mandatory']]
                    question_sub_text = row[col_idx['question_sub_text']] if 'question_sub_text' in col_idx else None

                    # Skip rows without question text
                    if not question_text:
                        continue

                    # Create template if not exists (use first row's data)
                    if template is None:
                        if not survey_name or version is None:
                            self.message_user(
                                request,
                                'اسم الاستبيان والإصدار مطلوبان في الصف الأول',
                                messages.ERROR,
                            )
                            return TemplateResponse(
                                request,
                                'admin/import_template.html',
                                {**self.admin_site.each_context(request), 'opts': self.model._meta},
                            )

                        template = SurveyTemplate.objects.create(
                            name=survey_name,
                            version=int(version),
                            is_active=True,
                        )

                    # Get or create question category
                    question_category = None
                    if category_name:
                        question_category, created = QuestionCategory.objects.get_or_create(
                            name=category_name
                        )
                        if created:
                            categories_created += 1

                    # Validate question type
                    if question_type not in valid_question_types:
                        self.message_user(
                            request,
                            f'نوع السؤال غير صالح في الصف {row_num}: {question_type}. الأنواع المتاحة: {", ".join(valid_question_types)}',
                            messages.ERROR,
                        )
                        continue

                    # Parse is_mandatory
                    if isinstance(is_mandatory, bool):
                        mandatory = is_mandatory
                    elif isinstance(is_mandatory, str):
                        mandatory = is_mandatory.lower() in ['true', 'yes', '1', 'نعم']
                    else:
                        mandatory = bool(is_mandatory) if is_mandatory is not None else True

                    # Create question
                    question = SurveyQuestion.objects.create(
                        question_text=question_text,
                        question_sub_text=question_sub_text if question_sub_text else None,
                        question_type=question_type,
                        question_category=question_category,
                        is_mandatory=mandatory,
                    )
                    question.template.add(template)
                    questions_created += 1

                self.message_user(
                    request,
                    f'تم استيراد الاستبيان "{template.name}" بنجاح. '
                    f'الأسئلة: {questions_created}، التصنيفات الجديدة: {categories_created}',
                    messages.SUCCESS,
                )
                from django.urls import reverse
                from django.http import HttpResponseRedirect
                return HttpResponseRedirect(reverse('admin:hita_evaluation_surveytemplate_changelist'))

            except Exception as e:
                self.message_user(
                    request,
                    f'خطأ في استيراد الملف: {str(e)}',
                    messages.ERROR,
                )

        context = {
            **self.admin_site.each_context(request),
            'title': 'استيراد استبيان - Import Survey Template',
            'opts': self.model._meta,
            'question_types': [(qt.value, qt.label) for qt in QuestionType],
        }

        return TemplateResponse(
            request,
            'admin/import_template.html',
            context,
        )


@admin.register(QuestionCategory)
class QuestionCategoryAdmin(admin.ModelAdmin):
    list_display = ['name']
    search_fields = ['name']


@admin.register(SurveyQuestion)
class SurveyQuestionAdmin(admin.ModelAdmin):
    list_display = ['question_text', 'question_type', 'question_category']
    list_filter = ['question_type', 'question_category', 'template']
    search_fields = ['question_text']
    filter_horizontal = ['template']


@admin.register(SurveySession)
class SurveySessionAdmin(admin.ModelAdmin):
    list_display = [
        'session_id',
        'status',
        'department',
        'regulation',
        'is_parallel',
        'created_at',
        'submitted_at',
    ]
    list_filter = ['status', 'department', 'regulation', 'is_parallel']
    readonly_fields = ['session_id', 'created_at', 'submitted_at']


def generate_term_report(courses, filename_prefix='term_report'):
    """Generate a single-sheet Excel report using separate simple queries joined in Python."""
    dept_labels = dict(Department.choices)
    category_labels = dict(Category.choices)
    semester_type_labels = dict(SemesterType.choices)
    grade_labels = dict(ProfessorGrade.choices)
    question_type_labels = dict(QuestionType.choices)

    course_ids = list(courses.values_list('id', flat=True))

    # --- Simple, flat lookups (no JOINs, no ORDER BY) ---
    courses_map = {
        c['id']: c for c in Course.objects.filter(id__in=course_ids).values(
            'id', 'subject_id', 'semester_id', 'regulations_id'
        )
    }
    subject_ids = {c['subject_id'] for c in courses_map.values()}
    semester_ids = {c['semester_id'] for c in courses_map.values()}
    course_reg_ids = {c['regulations_id'] for c in courses_map.values() if c['regulations_id']}

    subjects_map = {s['id']: s for s in Subject.objects.filter(id__in=subject_ids).values(
        'id', 'name', 'department', 'category', 'credit_hours'
    )}
    semesters_map = {s['id']: s for s in Semester.objects.filter(id__in=semester_ids).values(
        'id', 'year', 'type'
    )}
    cp_map = {
        (cp['course_id'], cp['professor_id']): grade_labels.get(cp['grade'], '-') if cp['grade'] else '-'
        for cp in CourseProfessor.objects.filter(course_id__in=course_ids).values('course_id', 'professor_id', 'grade')
    }

    # Fetch answers (no JOINs, no ORDER BY)
    raw_answers = list(SurveyAnswer.objects.filter(course_id__in=course_ids).values(
        'course_id', 'professor_id', 'question_id', 'survey_session_id',
        'yes_no_answer', 'rating_answer', 'text_answer',
    ))

    if not raw_answers:
        wb = Workbook()
        ws = wb.active
        ws.append(['لا توجد تقييمات'])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(output.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.xlsx"'
        return response

    session_ids = {a['survey_session_id'] for a in raw_answers}
    professor_ids = {a['professor_id'] for a in raw_answers}
    question_ids = {a['question_id'] for a in raw_answers}

    sessions_map = {s['id']: s for s in SurveySession.objects.filter(id__in=session_ids).values(
        'id', 'regulation_id', 'is_parallel', 'created_at'
    )}
    professors_map = {p['id']: p['full_name'] for p in Professor.objects.filter(id__in=professor_ids).values('id', 'full_name')}
    questions_map = {q['id']: q for q in SurveyQuestion.objects.filter(id__in=question_ids).values(
        'id', 'question_text', 'question_type', 'question_category_id'
    )}

    cat_ids = {q['question_category_id'] for q in questions_map.values() if q['question_category_id']}
    cats_map = {c['id']: c['name'] for c in QuestionCategory.objects.filter(id__in=cat_ids).values('id', 'name')}

    all_reg_ids = course_reg_ids | {s['regulation_id'] for s in sessions_map.values() if s['regulation_id']}
    regs_map = {r['id']: r['name'] for r in Regulation.objects.filter(id__in=all_reg_ids).values('id', 'name')}

    # --- Build Excel ---
    wb = Workbook()
    ws = wb.active
    ws.title = 'تقرير الاستبيان'

    headers = [
        'الفصل الدراسي', 'اللائحة', 'القسم', 'المادة', 'التصنيف', 'الساعات',
        'الأستاذ', 'درجة الأستاذ', 'تصنيف السؤال', 'السؤال', 'نوع السؤال',
        'الإجابة', 'تاريخ التقييم', 'نوع التعليم',
    ]
    ws.append(headers)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for answer in raw_answers:
        course = courses_map.get(answer['course_id'], {})
        subject = subjects_map.get(course.get('subject_id'), {})
        semester = semesters_map.get(course.get('semester_id'), {})
        session = sessions_map.get(answer['survey_session_id'], {})
        question = questions_map.get(answer['question_id'], {})

        q_type = question.get('question_type', '')
        yn = answer['yes_no_answer']
        if q_type == 'YN':
            answer_value = 'نعم' if yn is True else 'لا' if yn is False else '-'
        elif q_type in ['R', 'S']:
            answer_value = str(answer['rating_answer']) if answer['rating_answer'] is not None else '-'
        elif q_type == 'T':
            answer_value = answer['text_answer'] or '-'
        else:
            answer_value = '-'

        sem_type = semester.get('type', '')
        reg_id = session.get('regulation_id') or course.get('regulations_id')
        created_at = session.get('created_at')

        ws.append([
            f"{semester.get('year', '')} - {semester_type_labels.get(sem_type, sem_type)}",
            regs_map.get(reg_id, '-'),
            dept_labels.get(subject.get('department', ''), '-'),
            subject.get('name', '-'),
            category_labels.get(subject.get('category', ''), '-'),
            subject.get('credit_hours', '-'),
            professors_map.get(answer['professor_id'], '-'),
            cp_map.get((answer['course_id'], answer['professor_id']), '-'),
            cats_map.get(question.get('question_category_id'), '-'),
            question.get('question_text', '-'),
            question_type_labels.get(q_type, q_type),
            answer_value,
            created_at.strftime('%Y-%m-%d %H:%M') if created_at else '-',
            'التعليم الموازي' if session.get('is_parallel') else 'الساعات المعتمدة',
        ])

    col_widths = [15, 15, 25, 25, 20, 10, 20, 18, 20, 50, 15, 15, 18, 18]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.sheet_view.rightToLeft = True

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename_prefix}.xlsx"'
    return response


@admin.register(SurveyAnswer)
class SurveyAnswerAdmin(admin.ModelAdmin):
    list_display = [
        'get_semester',
        'course',
        'professor',
        'question',
        'yes_no_answer',
        'rating_answer',
    ]
    list_filter = [
        'course__semester__year',
        'course__semester__type',
        'course__subject__department',
        'question__question_type',
    ]
    search_fields = ['course__subject__name', 'professor__full_name', 'question__question_text']
    list_select_related = ['course__semester', 'course__subject', 'professor', 'question']
    change_list_template = 'admin/survey_answer_changelist.html'

    def get_semester(self, obj):
        s = obj.course.semester
        return f"{s.year} - {s.get_type_display()}"
    get_semester.short_description = 'الفصل'
    get_semester.admin_order_field = 'course__semester__year'

    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path(
                'export-term-report/',
                self.admin_site.admin_view(self.export_term_report_view),
                name='hita_evaluation_surveyanswer_export_term',
            ),
        ]
        return custom_urls + urls

    def export_term_report_view(self, request):
        courses = Course.objects.all()
        return generate_term_report(courses, 'survey_answers_report')
